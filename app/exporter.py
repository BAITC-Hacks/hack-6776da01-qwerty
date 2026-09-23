from io import BytesIO

from docx import Document


def make_docx(transcript: str, summary: str, tasks: list[dict[str, str]]) -> bytes:
    doc = Document()
    doc.add_heading("Протокол совещания", level=0)
    doc.add_heading("Саммари", level=1)
    doc.add_paragraph(summary)
    doc.add_heading("Поручения", level=1)
    table = doc.add_table(rows=1, cols=3)
    for cell, title in zip(table.rows[0].cells, ("Поручение", "Ответственный", "Срок")):
        cell.text = title
    for item in tasks:
        cells = table.add_row().cells
        cells[0].text = item["task"]
        cells[1].text = item["responsible"]
        cells[2].text = item["deadline"]
    doc.add_heading("Транскрипт", level=1)
    doc.add_paragraph(transcript)
    output = BytesIO()
    doc.save(output)
    return output.getvalue()
